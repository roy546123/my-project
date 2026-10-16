"""Read the user's workbook as data; do not evaluate formulas or execute content."""
import datetime as dt
import io
import re
import zipfile

import openpyxl


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def ticker(name):
    match = re.match(r"^([A-Z]{2,6})(?:\s*-|$)", name)
    return match[1] if match else name


def parse_workbook(content: bytes):
    # Bound decompressed size before openpyxl sees untrusted ZIP input.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(f.file_size for f in archive.infolist()) > 60_000_000:
            raise ValueError("압축 해제한 엑셀 파일이 너무 큽니다.")
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        if "보유종목(입력)" not in book.sheetnames:
            raise ValueError("'보유종목(입력)' 시트가 없습니다. 제공해 주신 투자관리 엑셀 형식을 사용해 주세요.")
        holdings = []
        for row in book["보유종목(입력)"].iter_rows(values_only=True):
            row = tuple(row) + (None,) * max(0, 14 - len(row))
            if not (numeric(row[1]) and row[2] and row[6] and numeric(row[7]) and numeric(row[8])):
                continue
            name = str(row[6])
            holdings.append(dict(broker=str(row[2]), account=str(row[3] or "일반계좌"),
                                 category=str(row[5] or "기타"), name=name, ticker=ticker(name),
                                 cost=round(row[7]), value=round(row[8]),
                                 quantity=float(row[11]) if numeric(row[11]) else None,
                                 note=str(row[13] or "")))
        if not holdings:
            raise ValueError("가져올 보유종목이 없습니다. 수식의 계산 결과가 저장된 엑셀인지 확인해 주세요.")
        snapshots = []
        if "일별 수익률" in book.sheetnames:
            for row in book["일별 수익률"].iter_rows(min_row=5, values_only=True):
                row = tuple(row) + (None,) * max(0, 9 - len(row))
                if isinstance(row[0], (dt.datetime, dt.date)) and numeric(row[1]) and numeric(row[2]):
                    snapshots.append(dict(date=row[0].strftime("%Y-%m-%d"), cost=round(row[1]),
                                          value=round(row[2]), note=str(row[8] or "")))
        snapshots.sort(key=lambda x: x["date"])
        # Infer DCA only from explicit per-day purchase notes, never narrative instructions.
        plans = []
        for i, h in enumerate(holdings):
            match = re.search(r"매일\s*([\d.]+)\s*만원", h["note"])
            if match:
                plans.append(dict(holding_index=i, amount=round(float(match[1]) * 10000)))
        warnings = []
        if snapshots and sum(h["value"] for h in holdings) != snapshots[-1]["value"]:
            warnings.append("최신 일별 평가금액과 보유종목 합계가 다릅니다. 원본의 두 값을 각각 보존합니다.")
        warnings.append("집행기록은 현재 보유종목에 이미 반영되어 있을 수 있어 중복 매수로 가져오지 않습니다.")
        warnings.append("실현손익·MDD·납입계획 시트는 이번 가져오기 범위에 포함되지 않습니다.")
        return dict(holdings=holdings, snapshots=snapshots, plans=plans, warnings=warnings,
                    date=snapshots[-1]["date"] if snapshots else dt.date.today().isoformat())
    finally:
        book.close()
