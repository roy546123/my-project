from contextlib import contextmanager
from datetime import datetime
from io import BytesIO
from pathlib import Path
import json
import mimetypes
import os
import secrets
import shutil
import sqlite3
import time
from zoneinfo import ZoneInfo

from fastapi import FastAPI, File, HTTPException, UploadFile, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, Field
import openpyxl

from .excel import parse_workbook
from .ocr import recognize, parse_text

ROOT = Path(__file__).resolve().parents[1]
DATA = Path(os.environ.get("INVEST_DATA_DIR", ROOT / ".data"))
DATA.mkdir(parents=True, exist_ok=True)
DB = DATA / "portfolio.sqlite3"
WEB = Path(os.environ.get("INVEST_WEB_DIR", ROOT / "dist"))
PREVIEWS = {}
# Windows registry entries can otherwise assign text/plain to .js files.
for suffix, content_type in {'.js': 'text/javascript', '.mjs': 'text/javascript', '.css': 'text/css', '.wasm': 'application/wasm'}.items():
    mimetypes.add_type(content_type, suffix)
app = FastAPI(title="모아 · 투자 기록", docs_url="/api/docs")


@contextmanager
def database():
    con = sqlite3.connect(DB, timeout=15)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


with database() as db:
    db.executescript('''
    CREATE TABLE IF NOT EXISTS holdings (
      id TEXT PRIMARY KEY, broker TEXT NOT NULL, account TEXT NOT NULL, category TEXT NOT NULL,
      name TEXT NOT NULL, ticker TEXT NOT NULL, cost INTEGER NOT NULL, value INTEGER NOT NULL,
      quantity REAL, note TEXT NOT NULL DEFAULT '', revision INTEGER NOT NULL DEFAULT 0);
    CREATE TABLE IF NOT EXISTS plans (
      id TEXT PRIMARY KEY, holding_id TEXT NOT NULL REFERENCES holdings(id), amount INTEGER NOT NULL,
      active INTEGER NOT NULL DEFAULT 1);
    CREATE TABLE IF NOT EXISTS trades (
      id TEXT PRIMARY KEY, date TEXT NOT NULL, holding_id TEXT NOT NULL REFERENCES holdings(id),
      amount INTEGER NOT NULL, quantity REAL, note TEXT NOT NULL DEFAULT '',
      source TEXT NOT NULL DEFAULT 'manual', plan_id TEXT, holding_revision INTEGER NOT NULL,
      created_at TEXT NOT NULL, previous_quantity REAL, UNIQUE(date, plan_id));
    CREATE TABLE IF NOT EXISTS snapshots (
      date TEXT PRIMARY KEY, cost INTEGER NOT NULL, value INTEGER NOT NULL, note TEXT NOT NULL DEFAULT '');
    CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
    ''')
    if 'previous_quantity' not in {r[1] for r in db.execute('PRAGMA table_info(trades)')}:
        db.execute('ALTER TABLE trades ADD COLUMN previous_quantity REAL')


def uid():
    return secrets.token_hex(12)


def today():
    return datetime.now(ZoneInfo("Asia/Seoul")).date().isoformat()


def check_date(value):
    try:
        parsed = datetime.strptime(value, "%Y-%m-%d").date()
        if parsed.isoformat() != value:
            raise ValueError()
    except ValueError:
        raise HTTPException(422, "날짜는 YYYY-MM-DD 형식이어야 합니다.")
    if value > today():
        raise HTTPException(422, "실제 기록은 미래 날짜로 저장할 수 없습니다.")


class Holding(BaseModel):
    broker: str = Field(min_length=1, max_length=50)
    account: str = Field(min_length=1, max_length=50)
    category: str = Field(default="기타", max_length=50)
    name: str = Field(min_length=1, max_length=150)
    ticker: str = Field(min_length=1, max_length=150)
    cost: int = Field(ge=0, le=10**14)
    value: int = Field(ge=0, le=10**14)
    quantity: float | None = Field(default=None, ge=0, le=10**12, allow_inf_nan=False)
    note: str = Field(default="", max_length=2000)


class Snapshot(BaseModel):
    date: str
    cost: int = Field(ge=0, le=10**14)
    value: int = Field(ge=0, le=10**14)
    note: str = Field(default="", max_length=2000)
    overwrite: bool = False


class Trade(BaseModel):
    date: str
    holding_id: str
    amount: int = Field(gt=0, le=10**12)
    quantity: float | None = Field(default=None, gt=0, le=10**10, allow_inf_nan=False)
    note: str = Field(default="", max_length=2000)
    plan_id: str | None = None
    source: str = Field(default="manual", pattern="^(manual|dca|screenshot)$")


class TradeBatch(BaseModel):
    trades: list[Trade] = Field(min_length=1, max_length=100)


class Plan(BaseModel):
    holding_id: str
    amount: int = Field(gt=0, le=10**10)
    active: bool = True


class ImportConfirm(BaseModel):
    token: str
    replace: bool = False


def backup():
    folder = DATA / "backups"
    folder.mkdir(exist_ok=True)
    path = folder / f"portfolio-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uid()[:6]}.sqlite3"
    with sqlite3.connect(DB) as src, sqlite3.connect(path) as dest:
        src.backup(dest)
    return path.name


def commit_import(parsed, filename, replace=False):
    with database() as db:
        has_data = db.execute("SELECT COUNT(*) FROM holdings").fetchone()[0] > 0
    if has_data and not replace:
        raise HTTPException(409, "기존 데이터가 있습니다. 교체 확인이 필요합니다.")
    backup_name = backup() if has_data else None
    with database() as db:
        for table in ["trades", "plans", "holdings", "snapshots"]:
            db.execute(f"DELETE FROM {table}")
        ids = []
        for h in parsed["holdings"]:
            hid = uid()
            ids.append(hid)
            fields = list(h)
            db.execute(f"INSERT INTO holdings (id,{','.join(fields)}) VALUES ({','.join('?' for _ in range(len(fields)+1))})", [hid] + list(h.values()))
        for p in parsed["plans"]:
            db.execute("INSERT INTO plans VALUES (?,?,?,1)", (uid(), ids[p["holding_index"]], p["amount"]))
        for s in parsed["snapshots"]:
            db.execute("INSERT OR REPLACE INTO snapshots VALUES (?,?,?,?)", (s["date"], s["cost"], s["value"], s["note"]))
        db.execute("INSERT OR REPLACE INTO meta VALUES ('import',?)", (json.dumps({"filename": filename, "date": parsed["date"], "warnings": parsed["warnings"]}, ensure_ascii=False),))
    return {"ok": True, "backup": backup_name}


@app.get("/api/state")
def state():
    with database() as db:
        holdings = [dict(r) for r in db.execute("SELECT * FROM holdings ORDER BY rowid")]
        plans = [dict(r) for r in db.execute("SELECT * FROM plans ORDER BY rowid")]
        trades = [dict(r) for r in db.execute("SELECT * FROM trades ORDER BY date DESC,created_at DESC")]
        snaps = [dict(r) for r in db.execute("SELECT * FROM snapshots ORDER BY date")]
        meta = {r["key"]: json.loads(r["value"]) for r in db.execute("SELECT * FROM meta")}
    previous = None
    for s in snaps:
        s["profit"] = s["value"] - s["cost"]
        s["market_change"] = s["value"] - previous["value"] - (s["cost"] - previous["cost"]) if previous else None
        s["market_rate"] = s["market_change"] / previous["value"] if previous and previous["value"] else None
        previous = s
    return dict(holdings=holdings, plans=plans, trades=trades, snapshots=snaps, meta=meta, today=today(),
                ocr_engine='server' if shutil.which('tesseract') else 'browser')


@app.get('/api/health')
def health():
    import hashlib
    result = {'application': 'moa-invest', 'profile': hashlib.sha256(str(DATA.resolve()).encode()).hexdigest()[:16]}
    if getattr(app.state, 'desktop_bundle', None):
        result['bundle'] = app.state.desktop_bundle
    return result


@app.post('/api/desktop/stop')
async def stop_desktop(request: Request):
    expected = getattr(app.state, 'desktop_stop_token', None)
    stop = getattr(app.state, 'desktop_stop', None)
    if not expected or not stop:
        raise HTTPException(404, '종료할 데스크톱 서버가 없습니다.')
    supplied = request.headers.get('X-Moa-Desktop-Token', '')
    if not secrets.compare_digest(supplied.encode(), expected.encode()):
        raise HTTPException(403, '이 실행 세션의 종료 권한이 없습니다.')
    import asyncio
    asyncio.get_running_loop().call_later(.2, stop)
    return {'ok': True}


@app.post("/api/holdings")
def add_holding(h: Holding):
    with database() as db:
        values = h.model_dump()
        hid = uid()
        db.execute(f"INSERT INTO holdings (id,{','.join(values)}) VALUES ({','.join('?' for _ in range(len(values)+1))})", [hid] + list(values.values()))
    return {"id": hid}


@app.put("/api/holdings/{hid}")
def update_holding(hid: str, h: Holding):
    with database() as db:
        values = h.model_dump()
        cursor = db.execute(f"UPDATE holdings SET {','.join(f'{k}=?' for k in values)},revision=revision+1 WHERE id=?", list(values.values()) + [hid])
        if not cursor.rowcount:
            raise HTTPException(404, "보유종목을 찾을 수 없습니다.")
    return {"ok": True}


@app.post("/api/snapshots")
def save_snapshot(s: Snapshot):
    check_date(s.date)
    with database() as db:
        if not s.overwrite and db.execute("SELECT 1 FROM snapshots WHERE date=?", (s.date,)).fetchone():
            raise HTTPException(409, "해당 날짜의 기록이 있습니다. 기존 기록 수정을 선택해 주세요.")
        db.execute("INSERT OR REPLACE INTO snapshots VALUES (?,?,?,?)", (s.date, s.cost, s.value, s.note))
    return {"ok": True}


@app.delete("/api/snapshots/{date}")
def delete_snapshot(date: str):
    with database() as db:
        db.execute("DELETE FROM snapshots WHERE date=?", (date,))
    return {"ok": True}


@app.post("/api/trades")
def add_trades(batch: TradeBatch):
    with database() as db:
        db.execute("BEGIN IMMEDIATE")
        for t in batch.trades:
            check_date(t.date)
            if t.source == 'screenshot' and not t.plan_id:
                match = db.execute('SELECT id FROM plans WHERE holding_id=? AND active=1', (t.holding_id,)).fetchone()
                if match:
                    t = t.model_copy(update={'plan_id': match['id']})
            h = db.execute("SELECT * FROM holdings WHERE id=?", (t.holding_id,)).fetchone()
            if not h:
                raise HTTPException(404, "매수 대상 종목을 찾을 수 없습니다.")
            if t.plan_id:
                plan = db.execute("SELECT * FROM plans WHERE id=? AND holding_id=? AND active=1", (t.plan_id, t.holding_id)).fetchone()
                if not plan:
                    raise HTTPException(422, "유효한 DCA 계획이 아닙니다.")
            # Manual/screenshot duplicate detection is conservative; override by entering a distinct memo.
            if not t.plan_id and db.execute("SELECT 1 FROM trades WHERE date=? AND holding_id=? AND amount=? AND note=?", (t.date,t.holding_id,t.amount,t.note)).fetchone():
                raise HTTPException(409, "동일한 매수 기록이 있습니다. 별도 체결이라면 메모를 달리해 주세요.")
            rev = h["revision"] + 1
            quantity = h["quantity"] + t.quantity if h["quantity"] is not None and t.quantity is not None else None
            try:
                db.execute("INSERT INTO trades VALUES (?,?,?,?,?,?,?,?,?,?,?)", (uid(),t.date,t.holding_id,t.amount,t.quantity,t.note,t.source,t.plan_id,rev,datetime.now().isoformat(),h['quantity']))
            except sqlite3.IntegrityError:
                raise HTTPException(409, "해당 날짜에 이미 기록한 DCA 매수가 있습니다. 중복 저장하지 않았습니다.")
            db.execute("UPDATE holdings SET cost=cost+?,value=value+?,quantity=?,revision=? WHERE id=?", (t.amount,t.amount,quantity,rev,t.holding_id))
    return {"ok": True, "count": len(batch.trades)}


@app.delete("/api/trades/{tid}")
def cancel_trade(tid: str):
    with database() as db:
        db.execute("BEGIN IMMEDIATE")
        t = db.execute("SELECT * FROM trades WHERE id=?", (tid,)).fetchone()
        if not t:
            raise HTTPException(404, "기록이 없습니다.")
        h = db.execute("SELECT * FROM holdings WHERE id=?", (t["holding_id"],)).fetchone()
        if h["revision"] != t["holding_revision"]:
            raise HTTPException(409, "이후 매수나 평가금액 수정이 있어 취소할 수 없습니다. 해당 종목의 최신 기록부터 취소해 주세요. 평가금액을 직접 수정한 경우 보유종목에서 정정해 주세요.")
        quantity = t['previous_quantity']
        db.execute("UPDATE holdings SET cost=cost-?,value=value-?,quantity=?,revision=revision-1 WHERE id=?", (t["amount"],t["amount"],quantity,t["holding_id"]))
        db.execute("DELETE FROM trades WHERE id=?", (tid,))
    return {"ok": True}


@app.post("/api/plans")
def add_plan(p: Plan):
    with database() as db:
        if not db.execute("SELECT 1 FROM holdings WHERE id=?", (p.holding_id,)).fetchone():
            raise HTTPException(404, "보유종목을 먼저 등록해 주세요.")
        if db.execute("SELECT 1 FROM plans WHERE holding_id=? AND active=1",(p.holding_id,)).fetchone():
            raise HTTPException(409, "이 종목은 이미 DCA 계획에 있습니다.")
        db.execute("INSERT INTO plans VALUES (?,?,?,?)", (uid(),p.holding_id,p.amount,int(p.active)))
    return {"ok": True}


@app.put("/api/plans/{pid}")
def update_plan(pid: str, p: Plan):
    with database() as db:
        if not db.execute("SELECT 1 FROM plans WHERE id=? AND holding_id=?", (pid,p.holding_id)).fetchone():
            raise HTTPException(404, "DCA 계획이 없습니다.")
        db.execute("UPDATE plans SET amount=?,active=? WHERE id=?", (p.amount,int(p.active),pid))
    return {"ok": True}


async def read_upload(file, max_size=10_000_000):
    content = await file.read(max_size + 1)
    if len(content) > max_size:
        raise HTTPException(413, "파일은 10MB 이하로 올려 주세요.")
    return content


@app.post("/api/import/preview")
async def preview(file: UploadFile = File(...)):
    content = await read_upload(file)
    try:
        parsed = await run_in_threadpool(parse_workbook, content)
    except Exception as e:
        raise HTTPException(422, str(e) if isinstance(e, ValueError) else "엑셀을 읽을 수 없습니다. 정상적인 .xlsx 파일인지 확인해 주세요.")
    for k in list(PREVIEWS):
        if time.time() - PREVIEWS[k][0] > 900:
            del PREVIEWS[k]
    if len(PREVIEWS) >= 10:
        del PREVIEWS[next(iter(PREVIEWS))]
    token = uid()
    PREVIEWS[token] = (time.time(), parsed, Path(file.filename or "투자관리.xlsx").name)
    return {"token": token, **parsed}


@app.post("/api/import/confirm")
def confirm_import(data: ImportConfirm):
    entry = PREVIEWS.get(data.token)
    if not entry or time.time() - entry[0] > 900:
        raise HTTPException(410, "미리보기가 만료되었습니다. 파일을 다시 선택해 주세요.")
    result = commit_import(entry[1],entry[2],data.replace)
    PREVIEWS.pop(data.token, None)
    return result


@app.post("/api/ocr")
async def ocr(file: UploadFile = File(...)):
    content = await read_upload(file)
    try:
        return await run_in_threadpool(recognize, content, DATA / "tessdata")
    except (ValueError, OSError) as e:
        raise HTTPException(422, str(e) if isinstance(e, ValueError) else "이미지 형식 또는 OCR 설치 상태를 확인해 주세요. PNG·JPG·WEBP를 지원합니다.")
    except Exception:
        raise HTTPException(422, "이미지 인식이 완료되지 않았습니다. 더 작은 이미지로 다시 시도해 주세요.")


class OcrText(BaseModel):
    text: str = Field(min_length=1, max_length=80000)


@app.post('/api/ocr/parse')
def parse_browser_ocr(data: OcrText):
    return {'text': data.text, 'candidates': parse_text(data.text)}


@app.get("/api/export.xlsx")
def export_excel():
    data = state()
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "보유종목(입력)"
    sheet.append([None,"NO","투자처","계정","계좌 구분","자산 구분","상품명","투자원금","평가금액","수익금액","수익률","주식수","평균매입가","비고"])
    for i,h in enumerate(data["holdings"],1):
        sheet.append([None,i,h["broker"],h["account"],h["broker"]+" - "+h["account"],h["category"],h["name"],h["cost"],h["value"],h["value"]-h["cost"],(h["value"]/h["cost"]-1) if h["cost"] else None,h["quantity"],None,h["note"]])
    sheet = book.create_sheet("일별 수익률")
    for row in [["모아 · 자산 기록"],[],[],["날짜","투자원금","평가금액","원금 증감","평가손익","수익률","시장 손익","시장 수익률","메모"]]:
        sheet.append(row)
    prev = None
    for s in data["snapshots"]:
        sheet.append([datetime.fromisoformat(s["date"]),s["cost"],s["value"],s["cost"]-prev["cost"] if prev else None,s["profit"],s["profit"]/s["cost"] if s["cost"] else None,s["market_change"],s["market_rate"],s["note"]])
        prev=s
    for key,title in [("trades","매수기록"),("plans","DCA계획")]:
        sheet = book.create_sheet(title)
        if data[key]:
            fields=list(data[key][0])
            sheet.append(fields)
            for row in data[key]:
                sheet.append([row[k] for k in fields])
    # Prevent formula injection by treating every user-controlled string as text.
    for s in book:
        s.freeze_panes="C2"
        for row in s:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type="s"
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":"attachment; filename=moa-portfolio.xlsx"})


@app.get("/api/backup")
def download_backup():
    name = backup()
    return FileResponse(DATA / "backups" / name, filename=name, media_type="application/vnd.sqlite3")


@app.middleware("http")
async def local_api_guard(request, call_next):
    # The app has no public login: same-origin writes only, bind to loopback by default.
    origin=request.headers.get("origin")
    if request.method not in ("GET","HEAD","OPTIONS") and origin:
        from urllib.parse import urlsplit
        if urlsplit(origin).netloc != request.headers.get("host"):
            return __import__('starlette.responses',fromlist=['JSONResponse']).JSONResponse({"detail":"다른 사이트에서의 데이터 변경은 허용하지 않습니다."},status_code=403)
    response=await call_next(request)
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"]="no-store"
    elif response.headers.get('content-type', '').startswith('text/html'):
        # Revalidate the HTML entry point so a rebuild cannot leave cached bundle URLs.
        response.headers['Cache-Control']='no-cache'
    response.headers["X-Content-Type-Options"]="nosniff"
    response.headers["X-Frame-Options"]="DENY"
    return response


if WEB.exists():
    app.mount("/",StaticFiles(directory=WEB,html=True),name="frontend")
