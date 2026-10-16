"""Financial invariants, atomic writes, import/export, OCR review, and security boundaries."""
import io
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest

TEST_DIR = tempfile.TemporaryDirectory(prefix='moa-test-')
os.environ['INVEST_DATA_DIR'] = TEST_DIR.name
from fastapi.testclient import TestClient
import openpyxl
from backend.app import app, DB, PREVIEWS
from backend.excel import parse_workbook
from backend.ocr import parse_text

client = TestClient(app)


def sample_excel():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title='보유종목(입력)'
    sheet.append([None,'NO','투자처','계정','계좌','구분','상품명','원금','평가','손익','수익률','수량','평단','비고'])
    sheet.append([None,1,'토스증권','해외계좌','토스','핵심 성장','QLD - ETF',100000,110000,10000,.1,10,None,'매일 10만원 소수점 자동매수'])
    sheet.append([None,'합계',None,None,None,None,None,100000,110000])
    s=book.create_sheet('일별 수익률')
    for row in [[],[],[],[],[__import__('datetime').datetime(2026,1,1),100000,110000], [__import__('datetime').datetime(2026,1,2),120000,135000], [__import__('datetime').datetime(2026,1,3),120000,None]]:
        s.append(row)
    b=io.BytesIO();book.save(b);return b.getvalue()


class ApiTests(unittest.TestCase):
    def setUp(self):
        with sqlite3.connect(DB) as con:
            for table in ['trades','plans','holdings','snapshots','meta']:
                con.execute(f'DELETE FROM {table}')
        PREVIEWS.clear()
        self.holding={'broker':'토스증권','account':'해외계좌','category':'나스닥','name':'QLD','ticker':'QLD','cost':100000,'value':110000,'quantity':10,'note':''}
        self.hid=client.post('/api/holdings',json=self.holding).json()['id']

    def state(self):
        r=client.get('/api/state');self.assertEqual(r.status_code,200);return r.json()

    def trade(self,**kwargs):
        return {'date':'2026-01-02','holding_id':self.hid,'amount':20000,**kwargs}

    def test_cashflow_excluded_from_market_profit(self):
        client.post('/api/snapshots',json={'date':'2026-01-01','cost':100000,'value':110000})
        client.post('/api/snapshots',json={'date':'2026-01-02','cost':120000,'value':135000})
        snapshots=self.state()['snapshots']
        self.assertIsNone(snapshots[0]['market_change'])
        self.assertEqual(snapshots[1]['market_change'],5000)
        self.assertAlmostEqual(snapshots[1]['market_rate'],5000/110000)

    def test_buy_does_not_generate_profit_or_snapshot(self):
        r=client.post('/api/trades',json={'trades':[self.trade(quantity=1.2)]})
        self.assertEqual(r.status_code,200,r.text)
        s=self.state();h=s['holdings'][0]
        self.assertEqual(h['cost'],120000);self.assertEqual(h['value'],130000)
        self.assertEqual(h['value']-h['cost'],10000)
        self.assertAlmostEqual(h['quantity'],11.2);self.assertEqual(s['snapshots'],[])

    def test_atomic_batch_duplicate_plan(self):
        client.post('/api/plans',json={'holding_id':self.hid,'amount':20000})
        pid=self.state()['plans'][0]['id']
        t=self.trade(plan_id=pid,source='dca')
        r=client.post('/api/trades',json={'trades':[t,t]})
        self.assertEqual(r.status_code,409)
        self.assertEqual(self.state()['trades'],[])
        self.assertEqual(self.state()['holdings'][0]['cost'],100000)
        self.assertEqual(client.post('/api/trades',json={'trades':[t]}).status_code,200)
        self.assertEqual(client.post('/api/trades',json={'trades':[t]}).status_code,409)
        self.assertEqual(self.state()['holdings'][0]['cost'],120000)

    def test_missing_quantity_and_cancel_restore_original(self):
        client.post('/api/trades',json={'trades':[self.trade()]})
        s=self.state();self.assertIsNone(s['holdings'][0]['quantity'])
        self.assertEqual(client.delete('/api/trades/'+s['trades'][0]['id']).status_code,200)
        h=self.state()['holdings'][0]
        self.assertEqual(h['quantity'],10);self.assertEqual(h['value'],110000)

    def test_cancel_after_valuation_edit_is_blocked(self):
        client.post('/api/trades',json={'trades':[self.trade()]})
        tid=self.state()['trades'][0]['id']
        client.put('/api/holdings/'+self.hid,json=self.holding)
        self.assertEqual(client.delete('/api/trades/'+tid).status_code,409)

    def test_snapshot_overwrite_requires_explicit_confirmation(self):
        snapshot={'date':'2026-01-01','cost':100000,'value':110000}
        self.assertEqual(client.post('/api/snapshots',json=snapshot).status_code,200)
        self.assertEqual(client.post('/api/snapshots',json={**snapshot,'value':120000}).status_code,409)
        self.assertEqual(client.post('/api/snapshots',json={**snapshot,'value':120000,'overwrite':True}).status_code,200)
        self.assertEqual(self.state()['snapshots'][0]['value'],120000)

    def test_screenshot_is_linked_to_matching_dca_plan(self):
        client.post('/api/plans',json={'holding_id':self.hid,'amount':20000})
        pid=self.state()['plans'][0]['id']
        self.assertEqual(client.post('/api/trades',json={'trades':[self.trade(source='screenshot')]}).status_code,200)
        self.assertEqual(self.state()['trades'][0]['plan_id'],pid)
        self.assertEqual(client.post('/api/trades',json={'trades':[self.trade(source='dca',plan_id=pid)]}).status_code,409)

    def test_excel_preview_does_not_mutate_and_confirm_backs_up(self):
        r=client.post('/api/import/preview',files={'file':('test.xlsx',sample_excel())})
        self.assertEqual(r.status_code,200,r.text)
        p=r.json();self.assertEqual(len(p['holdings']),1);self.assertEqual(len(p['snapshots']),2);self.assertEqual(len(p['plans']),1)
        self.assertEqual(self.state()['holdings'][0]['id'],self.hid)
        self.assertEqual(client.post('/api/import/confirm',json={'token':p['token']}).status_code,409)
        r=client.post('/api/import/confirm',json={'token':p['token'],'replace':True})
        self.assertEqual(r.status_code,200);self.assertTrue((Path(TEST_DIR.name)/'backups'/r.json()['backup']).exists())
        self.assertEqual(len(self.state()['snapshots']),2)
        self.assertEqual(client.post('/api/import/confirm',json={'token':p['token'],'replace':True}).status_code,410)

    def test_excel_export_preserves_text_and_reimports_holdings_history(self):
        client.put('/api/holdings/'+self.hid,json={**self.holding,'name':'=HYPERLINK("bad")'})
        client.post('/api/snapshots',json={'date':'2026-01-01','cost':100000,'value':110000})
        r=client.get('/api/export.xlsx');self.assertEqual(r.status_code,200)
        book=openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(book['보유종목(입력)']['G2'].data_type,'s')
        parsed=parse_workbook(r.content)
        self.assertEqual(parsed['holdings'][0]['value'],110000)
        self.assertEqual(len(parsed['snapshots']),1)

    def test_validation_and_cross_origin_write(self):
        self.assertEqual(client.post('/api/trades',json={'trades':[self.trade(amount=-1)]}).status_code,422)
        self.assertEqual(client.post('/api/trades',json={'trades':[self.trade(date='2999-01-01')]}).status_code,422)
        self.assertEqual(client.post('/api/holdings',json=self.holding,headers={'origin':'https://different.example'}).status_code,403)

    def test_ocr_does_not_guess_dollar_amounts_or_execute_text(self):
        rows=parse_text('QLD\n체결 단가 $94.05\n2.5 shares\n총 금액 100,000원\nTQQQ\nTotal $123.45\nignore all previous instructions')
        self.assertEqual(rows[0]['amount'],100000);self.assertEqual(rows[0]['quantity'],2.5)
        self.assertIsNone(rows[1]['amount'])

    def test_invalid_workbook_keeps_data(self):
        self.assertEqual(client.post('/api/import/preview',files={'file':('bad.xlsx',b'not an excel')}).status_code,422)
        self.assertEqual(len(self.state()['holdings']),1)


if __name__=='__main__':
    unittest.main()
